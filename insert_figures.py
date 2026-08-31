# =============================================================================
# Citation notice
#
# This code accompanies the manuscript:
#   J. Tang, S. Zhang, S. Ran, H. Ho. "A Coupled Reduced-Order Model of
#   First-Trimester Fetal Ductus Arteriosus Flow." Manuscript in preparation
#   / under review as of the time this repository was published -- please
#   check for the final published citation (journal, year, volume, DOI) and
#   cite that version if available; otherwise cite this repository directly.
#
# If you use, adapt, or build on this code, please cite the paper above.
# =============================================================================

"""Insert the composed Graphical Abstract and model schematic into the
manuscript, and refresh the (previously buggy-labeled) embedded waveform
figure with its corrected version. Renumbers figures: schematic becomes
Figure 1 (Methods 2.2), waveform becomes Figure 2 (was 1), timing-offset
becomes Figure 3 (was 2)."""
import docx
from docx.shared import Inches

PATH = "paper/A Coupled Reduced-Order Model of Fetal Ductus Arteriosus Flow.docx"
d = docx.Document(PATH)
paras = d.paragraphs

# --- 1. Graphical Abstract: replace placeholder text with the composed image ---
ga_placeholder = paras[7]
for run in list(ga_placeholder.runs):
    run.text = ""
ga_placeholder.add_run().add_picture("paper_figures/Graphical_Abstract.png", width=Inches(6.3))

# --- 2. Renumber existing figure captions: "Figure 1." -> "Figure 2.", "Figure 2." -> "Figure 3." ---
for p in paras:
    if p.text.startswith("Figure 2."):
        for r in p.runs:
            if r.text.startswith("Figure 2."):
                r.text = r.text.replace("Figure 2.", "Figure 3.", 1)
                break
        else:
            p.runs[0].text = p.runs[0].text.replace("Figure 2.", "Figure 3.", 1)
for p in paras:
    if p.text.startswith("Figure 1."):
        for r in p.runs:
            if r.text.startswith("Figure 1."):
                r.text = r.text.replace("Figure 1.", "Figure 2.", 1)
                break
        else:
            p.runs[0].text = p.runs[0].text.replace("Figure 1.", "Figure 2.", 1)

# --- 3. Replace the outdated embedded waveform image (paragraph with a drawing, just
#     before the "Figure 2." -- was "Figure 1." -- caption) with the corrected PNG ---
for i, p in enumerate(paras):
    if p.text.startswith("Figure 2. Simulated fetal DA velocity waveform"):
        img_para = paras[i - 1]
        for run in list(img_para.runs):
            run.text = ""
            run._element.getparent().remove(run._element)
        img_para.add_run().add_picture("paper_figures/Figure_simulated_waveform.png", width=Inches(5.2))
        break

# --- 4. Insert the model schematic as Figure 1 in Methods 2.2 ---
for i, p in enumerate(paras):
    if p.text.startswith("2.2 Model structure"):
        anchor = paras[i + 1]  # first content paragraph of 2.2
        break

fig_para = anchor.insert_paragraph_before("")
fig_para.add_run().add_picture("paper_figures/Figure_model_schematic.png", width=Inches(5.6))
cap_para = anchor.insert_paragraph_before(
    "Figure 1. The coupled nonlinear-RL model: three simultaneously integrated states "
    "(pulmonary compartment pressure, aortic compartment pressure, ductal flow) with "
    "ductal compliance folded into the two compartments (Section 2.2)."
)
cap_para.runs[0].italic = True

d.save(PATH)
print("Saved with figures inserted/renumbered.")

# --- verify ---
d2 = docx.Document(PATH)
n_images = sum(1 for p in d2.paragraphs for r in p.runs if "graphic" in r._element.xml.lower())
print("total embedded images now:", n_images)
for p in d2.paragraphs:
    if p.text.startswith("Figure"):
        print(" -", p.text[:90])
